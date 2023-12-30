import openpyxl
import re
from fpdf import FPDF

excel_file = "M290.23-24.PA-Noten.xlsx"  # Update the path to your file

workbook = openpyxl.load_workbook(excel_file, data_only=True)
sheet = workbook.active
headers = sheet[1]
maxPoints = sheet[2]

def add_header(pdf, header_text):
    pdf.set_font("DejaVu", size=12)
    pdf.cell(0, 10, header_text, ln=True, align='C')

def normalize(comment_text):
    # Regex pattern to identify the standard disclaimer in threaded comments
    disclaimer_pattern = r"\[Threaded comment\][\s\S]*?https://go\.microsoft\.com/fwlink/\?linkid=870924"
    # Remove the standard disclaimer
    stripped_text = re.sub(disclaimer_pattern, "", comment_text)
    # Remove "Comment: " from the string
    stripped_text = stripped_text.replace("Comment:", "")
    stripped_text = stripped_text.replace("None", "")
    # Remove extra whitespaces
    stripped_text = re.sub(r'\s+', ' ', stripped_text).strip()
    return stripped_text

# Iterate over each row in the worksheet, starting from the third row
for row_index, row in enumerate(sheet.iter_rows(min_row=3), start=2):
    # Initialize a new PDF document
    pdf = FPDF()
    pdf.add_font('DejaVu', '', 'DejaVuSans.ttf', uni=True)
    pdf.add_font('DejaVu', 'B', 'DejaVuSans-Bold.ttf', uni=True)
    pdf.add_page()
    add_header(pdf, f"M290 R.Hatz 01/2024 Gruppe: {row[0].value}") 
    pdf.set_font("DejaVu", size=10)
    pdf_name = f"_report_{row[0].value}.pdf"
    
    counter=0
    # Add the headers and corresponding row data to the PDF
    for cell in row:
        counter = counter + 1
        if counter == 1:
            continue
        cell_value = str(cell.value) if cell.value else ''
        cell_value = normalize(cell_value)

        header = headers[counter-1].value if headers[counter-1].value else '' 
        
       
        header = normalize(str(header))  
        max_points = maxPoints[counter-1].value 
        #max_points = normalize(max_points)

        entry = f"{header}: {cell_value} / {max_points}"
        if header and re.search('[a-zA-Z]', header):
            pdf.set_font("DejaVu", 'B', 10)   
        else:
            pdf.set_font("DejaVu", '', 10)   
        pdf.multi_cell(0, 8, entry, align='L')

        # Check for comments and append if present
        if cell.comment:
            comment_text = normalize(cell.comment.text.strip())
            pdf.set_font("DejaVu", size=8)  # Smaller font for comments
            
            for part in comment_text.split("- "): 
                part = re.sub(r'\s+', ' ', part).strip()
                if (part): 
                    pdf.multi_cell(0, 4, f"- {part}", align='L')
                    pdf.ln(1)
                
            pdf.ln(1)
            pdf.set_font("DejaVu", size=10)  # Reset font size to normal
       
    # Save the PDF file with the designated name
    pdf.output(pdf_name)
